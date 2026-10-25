from __future__ import annotations

import json
import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx import Document

from rag_engine import ABSTENTION, IngestionError, OrganizationalMemoryEngine, parse_file


class FileParsingTests(unittest.TestCase):
    def test_text_parser_preserves_markdown_section(self) -> None:
        parsed = parse_file("notes.md", b"# Decision\n\nDatabase choice: PostgreSQL\n")
        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0].section, "Decision")
        self.assertIn("Database choice: PostgreSQL", parsed[0].text)

    def test_json_parser_preserves_author_and_timestamp(self) -> None:
        parsed = parse_file(
            "messages.json",
            json.dumps([{"user": "Platform Engineer", "ts": "2026-06-13", "text": "Approved."}]).encode(),
        )
        self.assertEqual(parsed[0].author, "Platform Engineer")
        self.assertEqual(parsed[0].recorded_at, "2026-06-13")

    def test_csv_parser_includes_headers_with_values(self) -> None:
        parsed = parse_file("decisions.csv", b"date,decision\n2026-06-13,PostgreSQL selected\n")
        self.assertIn("decision: PostgreSQL selected", parsed[0].text)

    def test_docx_parser_preserves_heading_as_section(self) -> None:
        document = Document()
        document.add_heading("Decision", level=1)
        document.add_paragraph("Database choice: PostgreSQL")
        content = BytesIO()
        document.save(content)
        parsed = parse_file("decision.docx", content.getvalue())
        self.assertEqual(parsed[0].section, "Decision")
        self.assertIn("Database choice: PostgreSQL", parsed[0].text)

    def test_pdf_parser_rejects_unreadable_pdf(self) -> None:
        with self.assertRaises(IngestionError):
            parse_file("broken.pdf", b"not a PDF")

    def test_rejects_unsupported_and_malformed_uploads(self) -> None:
        with self.assertRaises(IngestionError):
            parse_file("program.exe", b"anything")
        with self.assertRaises(IngestionError):
            parse_file("broken.json", b"{")


class RetrievalTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.model_env_patch = patch.dict(os.environ, {
            "OPENAI_API_KEY": "",
            "EMBEDDING_PROVIDER": "local",
        })
        self.model_env_patch.start()
        self.engine = OrganizationalMemoryEngine(self.temp_dir.name, min_score=0.12)
        self.engine.ingest_bytes(
            "adr.md",
            b"# Accepted decision\nDecision date: 2026-06-13\n\n"
            b"Database choice: PostgreSQL. The Payments Pilot selected PostgreSQL "
            b"for atomic ledger updates and weekly reconciliation.",
        )
        self.engine.ingest_bytes(
            "handover.md",
            b"# Handover\nUpdated: 2026-06-15\n\n"
            b"Database choice: MongoDB. This old handover line was not updated after "
            b"the accepted decision in adr.md.",
        )

    def tearDown(self) -> None:
        self.model_env_patch.stop()
        self.temp_dir.cleanup()

    def test_supported_question_returns_actual_source_citations(self) -> None:
        result = self.engine.query("Which database was selected for the Payments Pilot?")
        self.assertFalse(result["abstained"])
        self.assertTrue(result["citations"])
        self.assertTrue(any(c["filename"] == "adr.md" for c in result["citations"]))
        for citation in result["citations"]:
            self.assertTrue(citation["snippet"])
            self.assertIn(citation["filename"], {"adr.md", "handover.md"})
            self.assertTrue(citation["id"].startswith("S"))

    def test_unsupported_question_abstains(self) -> None:
        result = self.engine.query("What is the team's office address and phone number?")
        self.assertEqual(result["answer"], ABSTENTION)
        self.assertTrue(result["abstained"])
        self.assertEqual(result["citations"], [])

    def test_citation_metadata_maps_to_source_document(self) -> None:
        citation = self.engine.query("PostgreSQL ledger reconciliation")["citations"][0]
        document = next(item for item in self.engine.list_documents() if item["filename"] == citation["filename"])
        self.assertGreater(document["stored_chunks"], 0)
        self.assertIn("PostgreSQL", citation["snippet"])
        self.assertGreaterEqual(citation["score"], self.engine.min_score)

    def test_duplicate_ingestion_is_skipped_by_content_hash(self) -> None:
        content = b"# New doc\n\nA distinct new document with a unique fact."
        first = self.engine.ingest_bytes("unique.md", content)
        second = self.engine.ingest_bytes("renamed.md", content)
        self.assertEqual(first["status"], "indexed")
        self.assertEqual(second["status"], "duplicate")
        self.assertEqual(len(self.engine.list_documents()), 3)

    def test_conflict_detector_returns_both_explicit_values(self) -> None:
        report = self.engine.detect_conflicts("Payments Pilot database choice")
        conflict = next(item for item in report["conflicts"] if item["key"] == "database choice")
        self.assertEqual({item["value"] for item in conflict["statements"]},
                         {"postgresql", "mongodb"})
        self.assertEqual({item["filename"] for item in conflict["statements"]},
                         {"adr.md", "handover.md"})

    def test_handover_marks_undocumented_information(self) -> None:
        handover = self.engine.generate_handover("Payments Pilot")
        self.assertEqual(
            handover["sections"]["Documented contacts or roles"][0]["excerpt"],
            "Not found in the available documents.",
        )

    def test_model_errors_fall_back_to_cited_excerpts(self) -> None:
        with patch.dict(os.environ, {"OPENAI_API_KEY": "test-key"}), patch(
            "rag_engine._openai_request", side_effect=RuntimeError("offline")
        ):
            result = self.engine.query("Which database was selected for the Payments Pilot?")
        self.assertEqual(result["mode"], "evidence-only")
        self.assertIn("offline", result["notice"])
        self.assertTrue(result["citations"])

    def test_missing_model_key_uses_evidence_only_mode(self) -> None:
        with patch.dict(os.environ, {"OPENAI_API_KEY": ""}):
            result = self.engine.query("Which database was selected for the Payments Pilot?")
        self.assertEqual(result["mode"], "evidence-only")
        self.assertTrue(result["citations"])

    def test_library_metrics_reflect_persisted_records(self) -> None:
        metrics = self.engine.metrics()
        self.assertEqual(metrics["documents"], 2)
        self.assertGreater(metrics["chunks"], 0)
        self.assertEqual(metrics["queries"], 0)

    def test_index_directories_scans_only_the_selected_folder(self) -> None:
        isolated_data = Path(self.temp_dir.name) / "isolated"
        engine = OrganizationalMemoryEngine(isolated_data)
        demo = isolated_data / "demo"
        docs = isolated_data / "docs"
        demo.mkdir()
        docs.mkdir()
        (demo / "sample.md").write_text("Demo folder fact: sample only.", encoding="utf-8")
        (docs / "private.md").write_text("Unselected folder fact.", encoding="utf-8")
        report = engine.index_directories(("demo",))
        self.assertEqual(report["indexed"], 1)
        self.assertEqual([item["filename"] for item in engine.list_documents()], ["sample.md"])

    def test_demo_dataset_answers_and_abstains_as_documented(self) -> None:
        demo = Path(__file__).parent.parent / "data" / "demo"
        with tempfile.TemporaryDirectory() as temp_dir:
            engine = OrganizationalMemoryEngine(temp_dir)
            for source in sorted(demo.glob("*.md")):
                engine.ingest_file(source)
            answer = engine.query("Which database was selected for the Payments Pilot?")
            unsupported = engine.query("What is the approved annual compensation budget?")
            conflicts = engine.detect_conflicts("payments pilot database choice")
        self.assertFalse(answer["abstained"])
        self.assertTrue(any(item["filename"] == "03_adr_014.md" for item in answer["citations"]))
        self.assertEqual(unsupported["answer"], ABSTENTION)
        self.assertTrue(any(item["key"] == "database choice" for item in conflicts["conflicts"]))


if __name__ == "__main__":
    unittest.main()
